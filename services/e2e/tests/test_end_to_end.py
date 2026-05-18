import os
import time
import uuid
from datetime import date
from io import BytesIO

import httpx
from openpyxl import load_workbook


KEYCLOAK_URL = os.environ.get("KEYCLOAK_URL", "http://localhost:8080").rstrip("/")
NSI_URL = os.environ.get("NSI_URL", "http://localhost:8001").rstrip("/")
DOCS_URL = os.environ.get("DOCS_URL", "http://localhost:8002").rstrip("/")
ORIGIN = os.environ.get("ORIGIN", "http://localhost:5173")
RABBITMQ_MGMT_URL = os.environ.get("RABBITMQ_MGMT_URL", "http://rabbitmq:15672").rstrip("/")
RABBITMQ_USER = os.environ.get("RABBITMQ_DEFAULT_USER", "converter_mq")
RABBITMQ_PASS = os.environ.get("RABBITMQ_DEFAULT_PASS", "converter_mq_local")
SEED_SKUS = {
    "bulk": "BULK-CRUSH-M800-20-40-001",
    "bolt": "FAST-BOLT-20X60-DIN933-001",
    "nut": "FAST-NUT-M16-DIN934-001",
}


def wait_until(fn, timeout_s: int = 90, interval_s: float = 1.0, err: str = "timeout"):
    deadline = time.time() + timeout_s
    last_exc = None
    while time.time() < deadline:
        try:
            value = fn()
            if value:
                return value
        except Exception as exc:
            last_exc = exc
        time.sleep(interval_s)
    if last_exc:
        raise RuntimeError(f"{err}: last error: {last_exc}") from last_exc
    raise RuntimeError(err)


def wait_ok(url: str, timeout_s: int = 60) -> None:
    deadline = time.time() + timeout_s
    last = None
    while time.time() < deadline:
        try:
            r = httpx.get(url, timeout=5)
            if r.status_code == 200:
                return
            last = f"{r.status_code} {r.text[:200]}"
        except Exception as e:
            last = str(e)
        time.sleep(1)
    raise RuntimeError(f"Service not ready: {url}. Last={last}")


def get_token(username: str, password: str) -> str:
    def call():
        response = httpx.post(
            f"{KEYCLOAK_URL}/realms/uom/protocol/openid-connect/token",
            data={
                "grant_type": "password",
                "client_id": "uom-cli",
                "username": username,
                "password": password,
            },
            headers={"Content-Type": "application/x-www-form-urlencoded"},
            timeout=15,
        )
        if response.status_code != 200:
            return None
        return response.json().get("access_token")

    return wait_until(call, err=f"failed to get token for {username}")


def auth_headers(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def select_seed_items(client: httpx.Client, token: str) -> dict[str, dict]:
    def call():
        response = client.get(f"{NSI_URL}/api/v1/items/", headers=auth_headers(token))
        response.raise_for_status()
        by_sku = {str(x.get("sku")): x for x in response.json()}
        if not all(sku in by_sku for sku in SEED_SKUS.values()):
            return None
        return {
            alias: by_sku[sku]
            for alias, sku in SEED_SKUS.items()
        }

    return wait_until(call, err="seeded NSI items not ready")


def wait_celery_worker_ready(client: httpx.Client) -> None:
    def call():
        response = client.get(
            f"{RABBITMQ_MGMT_URL}/api/queues/%2F/celery",
            auth=(RABBITMQ_USER, RABBITMQ_PASS),
        )
        if response.status_code == 404:
            return None
        response.raise_for_status()
        queue = response.json()
        return queue if int(queue.get("consumers") or 0) > 0 else None

    wait_until(call, err="celery worker consumer is not ready")


def wait_invoice_status(client: httpx.Client, token: str, invoice_id: int, want: str, timeout_s: int = 180) -> dict:
    last_status = None

    def call():
        nonlocal last_status
        response = client.get(f"{DOCS_URL}/api/v1/invoices/{invoice_id}/", headers=auth_headers(token))
        response.raise_for_status()
        invoice = response.json()
        status = invoice["status"]
        last_status = status
        if status == "failed":
            raise RuntimeError(f"invoice failed: {invoice.get('error')}")
        return invoice if status == want else None

    return wait_until(
        call,
        timeout_s=timeout_s,
        err=f"invoice {invoice_id} did not reach {want} (last_status={last_status})",
    )


def ensure_uom(client: httpx.Client, token: str, code: str, name: str, category_id: int, factor: str, precision: int):
    r = client.get(f"{NSI_URL}/api/v1/uoms/", headers=auth_headers(token))
    r.raise_for_status()
    for u in r.json():
        if u["code"] == code:
            return u["id"]

    r = client.post(
        f"{NSI_URL}/api/v1/uoms/",
        headers={**auth_headers(token), "Content-Type": "application/json"},
        json={
            "code": code,
            "name": name,
            "category": category_id,
            "factor_to_base": factor,
            "precision": precision,
        },
    )
    r.raise_for_status()
    return r.json()["id"]




def ensure_item_category(client: httpx.Client, token: str, name: str, default_uom_id: int) -> int:
    r = client.get(f"{NSI_URL}/api/v1/item-categories/", headers=auth_headers(token))
    r.raise_for_status()
    for c in r.json():
        if c["name"] == name:
            return c["id"]

    r = client.post(
        f"{NSI_URL}/api/v1/item-categories/",
        headers={**auth_headers(token), "Content-Type": "application/json"},
        json={"name": name, "default_uom": default_uom_id, "is_active": True},
    )
    r.raise_for_status()
    return r.json()["id"]

def test_openapi_and_full_flow():
    # Wait for services
    wait_ok(f"{NSI_URL}/healthz")
    wait_ok(f"{DOCS_URL}/healthz")

    token = get_token("operator", "operator")

    # --- CORS + OpenAPI checks (NSI + Documents) ---
    for base in (NSI_URL, DOCS_URL):
        r = httpx.get(
            f"{base}/api/schema/",
            headers={**auth_headers(token), "Origin": ORIGIN, "Accept": "application/json"},
            timeout=10,
        )
        assert r.status_code == 200
        assert r.headers.get("access-control-allow-origin") == ORIGIN
        data = r.json()
        assert "openapi" in data
        assert "info" in data and "title" in data["info"]

    with httpx.Client(timeout=20) as client:
        # --- Use default seeded items (no E2E item creation) ---
        seed_items = select_seed_items(client, token)
        bulk_item = seed_items["bulk"]
        fast_bolt_item = seed_items["bolt"]
        fast_nut_item = seed_items["nut"]

        # --- Create invoice in Documents ---
        inv_no = f"INV-E2E-{uuid.uuid4().hex[:6]}"
        inv = client.post(
            f"{DOCS_URL}/api/v1/invoices/",
            headers={**auth_headers(token), "Content-Type": "application/json"},
            json={
                "number": inv_no,
                "supplier": "ACME",
                "doc_date": date.today().isoformat(),
                "lines": [
                    {"line_no": 1, "item_id": int(bulk_item["id"]), "qty": "1.5", "uom_code": "TON", "context": {"item_name": str(bulk_item["name"])}},
                    {"line_no": 2, "item_id": int(fast_bolt_item["id"]), "qty": "10", "uom_code": "PCS", "context": {"item_name": str(fast_bolt_item["name"])}},
                    {"line_no": 3, "item_id": int(fast_nut_item["id"]), "qty": "100", "uom_code": "PCS", "context": {"item_name": str(fast_nut_item["name"])}},
                ],
            },
        )
        inv.raise_for_status()
        inv_id = inv.json()["id"]

        # --- Calculate ---
        wait_celery_worker_ready(client)
        r = client.post(f"{DOCS_URL}/api/v1/invoices/{inv_id}/calculate/", headers=auth_headers(token))
        r.raise_for_status()

        j = wait_invoice_status(client, token, inv_id, "calculated")
        # verify conversions
        by_line = {ln["line_no"]: ln for ln in j["lines"]}
        assert by_line[1]["converted"]["posting_uom_code"] == "M3"
        assert float(by_line[1]["converted"]["posting_qty"]) == 1.06383
        assert by_line[2]["converted"]["posting_uom_code"] == "KG"
        assert float(by_line[2]["converted"]["posting_qty"]) == 2.44
        assert by_line[3]["converted"]["posting_uom_code"] == "KG"
        assert float(by_line[3]["converted"]["posting_qty"]) == 3.33

        # --- Generate outputs ---
        r = client.post(f"{DOCS_URL}/api/v1/invoices/{inv_id}/generate/", headers=auth_headers(token))
        r.raise_for_status()

        generated_invoice = wait_invoice_status(client, token, inv_id, "generated")
        files = generated_invoice["files"]
        assert len(files) >= 2

        xlsx = [f for f in files if f["file_type"] == "xlsx"][0]
        pdf = [f for f in files if f["file_type"] == "pdf"][0]
        xlsx_url = xlsx.get("download_url") or xlsx.get("presigned_url")
        pdf_url = pdf.get("download_url") or pdf.get("presigned_url")
        assert xlsx_url, "download_url or presigned_url for xlsx should not be empty"
        assert pdf_url, "download_url or presigned_url for pdf should not be empty"

        # --- Download via presigned URL (no auth header needed) ---
        x = client.get(xlsx_url, headers=auth_headers(token) if xlsx.get("download_url") else None)
        x.raise_for_status()
        # XLSX is a zip file => starts with PK
        assert x.content[:2] == b"PK"

        wb = load_workbook(filename=BytesIO(x.content))
        ws = wb.active
        assert ws["A1"].value == "Номер накладной"
        assert ws["B1"].value == inv_no

        p = client.get(pdf_url, headers=auth_headers(token) if pdf.get("download_url") else None)
        p.raise_for_status()
        assert p.content[:4] == b"%PDF"
